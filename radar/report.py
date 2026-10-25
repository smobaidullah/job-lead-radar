"""Write the ranked leads as Excel, JSON and a static HTML page."""
from __future__ import annotations

import html
import json
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .sources import Job

SOURCE_LINKS = {
    "Remote OK": "https://remoteok.com",
    "Remotive": "https://remotive.com",
    "Arbeitnow": "https://www.arbeitnow.com",
}


def to_dict(job: Job) -> dict:
    return {
        "source": job.source, "id": job.id, "title": job.title, "company": job.company,
        "url": job.url, "posted": job.posted.isoformat(), "location": job.location,
        "salary": job.salary, "score": job.score, "matched": job.matched, "new": job.is_new,
        "snippet": job.description[:280],
    }


def write_json(jobs: list[Job], path: Path) -> None:
    payload = {"generated": datetime.now(timezone.utc).isoformat(), "jobs": [to_dict(j) for j in jobs]}
    path.write_text(json.dumps(payload, indent=1), encoding="utf-8")


def write_excel(jobs: list[Job], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    headers = ["Score", "New", "Title", "Company", "Location", "Salary", "Posted", "Matched keywords", "Source", "Link"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(vertical="center")

    new_fill = PatternFill("solid", fgColor="E0F2F0")
    for job in jobs:
        ws.append([
            job.score, "NEW" if job.is_new else "", job.title, job.company, job.location, job.salary,
            job.posted.replace(tzinfo=None), ", ".join(job.matched), job.source, "Open",
        ])
        row = ws.max_row
        link = ws.cell(row=row, column=10)
        link.hyperlink = job.url
        link.font = Font(color="0563C1", underline="single")
        ws.cell(row=row, column=7).number_format = "dd mmm yyyy"
        if job.is_new:
            for c in ws[row]:
                c.fill = new_fill

    for i, width in enumerate([7, 6, 48, 26, 22, 18, 12, 40, 11, 8], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    by_source = wb.create_sheet("Summary")
    by_source.append(["Source", "Leads"])
    for name in SOURCE_LINKS:
        by_source.append([name, sum(1 for j in jobs if j.source == name)])
    by_source.append(["Total", len(jobs)])
    by_source.column_dimensions["A"].width = 14
    for cell in by_source[1]:
        cell.font = Font(bold=True)
    wb.save(path)


def write_html(jobs: list[Job], path: Path, generated: datetime | None = None) -> None:
    generated = generated or datetime.now(timezone.utc)
    data = json.dumps([to_dict(j) for j in jobs]).replace("</", "<\\/")
    credits = " · ".join(f'<a href="{u}" rel="follow">{html.escape(n)}</a>' for n, u in SOURCE_LINKS.items())
    template = (Path(__file__).parent / "template.html").read_text(encoding="utf-8")
    page = (template
            .replace("{{DATA}}", data)
            .replace("{{GENERATED}}", generated.strftime("%d %b %Y, %H:%M UTC"))
            .replace("{{COUNT}}", str(len(jobs)))
            .replace("{{NEW}}", str(sum(j.is_new for j in jobs)))
            .replace("{{CREDITS}}", credits))
    path.write_text(page, encoding="utf-8")
